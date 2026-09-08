from __future__ import annotations

import csv
import io
import uuid
from datetime import date, datetime, time, timezone
from decimal import Decimal
from pathlib import Path

import psycopg2.extras
from openpyxl import load_workbook

from app.db import get_db
from app.services.orders import check_possible_duplicate, create_order
from app.services.validation import validate_order_fields
from app.utils.normalize import normalize_name

# Candidate normalized (lowercase, no spaces) header names Zomato/Swiggy-style
# exports commonly use for each target field. Import mapping is still fully
# user-adjustable -- this is just a best-effort starting guess.
FIELD_ALIASES = {
    "order_date": ["date", "orderdate", "dateplaced", "placedon", "orderedon", "orderdatetime"],
    "order_time": ["time", "ordertime", "timeplaced"],
    "platform": ["platform", "app", "source", "orderedfrom", "platformname"],
    "restaurant": ["restaurant", "restaurantname", "vendor", "outlet", "merchant", "storename"],
    "cuisine": ["cuisine", "category", "foodtype", "cuisinetype"],
    "food_items": ["items", "fooditems", "item", "orderitems", "dishes", "itemname", "itemsordered"],
    "subtotal": ["subtotal", "itemtotal", "subtotalamount", "itemstotal"],
    "discount": ["discount", "offer", "discountamount", "couponapplied", "couponsavings"],
    "delivery_fee": ["deliveryfee", "deliverycharge", "deliverycharges", "shippingfee"],
    "platform_fee": ["platformfee", "servicefee", "conveniencefee", "packagingcharge", "packagingfee"],
    "tax": ["tax", "taxes", "gst", "taxamount"],
    "total_amount": [
        "total", "totalamount", "amount", "grandtotal", "billamount",
        "ordertotal", "netamount", "amountpaid", "finalamount",
    ],
    "notes": ["notes", "remarks", "comment", "comments"],
}

TARGET_FIELD_LABELS = {
    "order_date": "Order Date *",
    "order_time": "Order Time",
    "platform": "Platform *",
    "restaurant": "Restaurant *",
    "cuisine": "Cuisine",
    "food_items": "Food Items",
    "subtotal": "Subtotal",
    "discount": "Discount",
    "delivery_fee": "Delivery Fee",
    "platform_fee": "Platform Fee",
    "tax": "Taxes",
    "total_amount": "Total Amount *",
    "notes": "Notes",
}

def _xlsx_cell_to_json_safe(cell):
    """openpyxl returns native datetime/date/time/Decimal objects for
    formatted cells, none of which are JSON-serializable -- and the session
    payload gets stored as JSONB. Normalize at the source so every consumer
    downstream just sees plain strings/numbers, same as a CSV would produce."""
    if cell is None:
        return ""
    if isinstance(cell, (datetime, date, time)):
        return cell.isoformat()
    if isinstance(cell, Decimal):
        return float(cell)
    return cell


def read_upload(file_storage):
    """Parse an uploaded CSV/XLSX FileStorage into (headers, rows-as-dicts)."""
    filename = file_storage.filename or ""
    ext = Path(filename).suffix.lower()

    if ext == ".csv":
        raw = file_storage.read()
        text = raw.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    elif ext in (".xlsx", ".xls"):
        wb = load_workbook(io.BytesIO(file_storage.read()), data_only=True)
        ws = wb.active
        rows = [
            [_xlsx_cell_to_json_safe(cell) for cell in row]
            for row in ws.iter_rows(values_only=True)
        ]
    else:
        raise ValueError("Unsupported file type. Please upload a .csv or .xlsx file.")

    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValueError("The file appears to be empty.")

    headers = [str(h).strip() for h in rows[0]]
    data_rows = []
    for r in rows[1:]:
        row_dict = {headers[i]: (r[i] if i < len(r) else "") for i in range(len(headers))}
        data_rows.append(row_dict)
    return headers, data_rows


def auto_detect_mapping(headers: list[str]) -> dict:
    normalized_headers = {h: normalize_name(h).replace(" ", "") for h in headers}
    mapping = {}
    used_headers = set()

    # Pass 1: exact alias matches, across ALL target fields, before any
    # fuzzy matching runs. Otherwise an early target's fuzzy substring check
    # (e.g. "subtotal" aliases) can steal a header that exactly matches a
    # later target (e.g. "Total Amount" belongs to total_amount, not subtotal).
    for target, aliases in FIELD_ALIASES.items():
        for h, norm in normalized_headers.items():
            if h in used_headers:
                continue
            if norm in aliases:
                mapping[target] = h
                used_headers.add(h)
                break

    # Pass 2: fuzzy substring matches for anything still unmapped.
    for target, aliases in FIELD_ALIASES.items():
        if target in mapping:
            continue
        for h, norm in normalized_headers.items():
            if h in used_headers or not norm:
                continue
            if any(alias in norm or norm in alias for alias in aliases):
                mapping[target] = h
                used_headers.add(h)
                break

    return mapping


def create_session(filename: str, headers: list[str], rows: list[dict], mapping: dict) -> str:
    session_id = uuid.uuid4().hex[:12]
    save_session(
        session_id,
        {
            "filename": filename,
            "headers": headers,
            "rows": rows,
            "mapping": mapping,
            "created_at": datetime.now(timezone.utc).isoformat(),
        },
    )
    return session_id


def load_session(session_id: str) -> dict | None:
    db = get_db()
    row = db.execute("SELECT payload FROM import_sessions WHERE id = ?", (session_id,)).fetchone()
    return row["payload"] if row else None


def save_session(session_id: str, payload: dict):
    db = get_db()
    db.execute(
        """
        INSERT INTO import_sessions (id, payload) VALUES (?, ?)
        ON CONFLICT (id) DO UPDATE SET payload = excluded.payload
        """,
        (session_id, psycopg2.extras.Json(payload)),
    )
    db.commit()


def delete_session(session_id: str):
    db = get_db()
    db.execute("DELETE FROM import_sessions WHERE id = ?", (session_id,))
    db.commit()


def apply_mapping(rows: list[dict], mapping: dict, fixed_platform: str | None = None) -> list[dict]:
    mapped = []
    for row in rows:
        item = {}
        for target, source_header in mapping.items():
            if source_header:
                item[target] = row.get(source_header, "")
        if fixed_platform and not str(item.get("platform") or "").strip():
            item["platform"] = fixed_platform
        mapped.append(item)
    return mapped


def validate_rows(mapped_rows: list[dict]) -> list[dict]:
    results = []
    seen_in_batch = set()
    for idx, raw in enumerate(mapped_rows, start=1):
        cleaned, errors = validate_order_fields(raw)
        is_duplicate = False
        if not errors:
            is_duplicate = check_possible_duplicate(
                cleaned["order_date"], cleaned["restaurant"], cleaned["total_amount"], cleaned["platform"]
            )
            batch_key = (
                cleaned["order_date"],
                normalize_name(cleaned["restaurant"]),
                cleaned["total_amount"],
                cleaned["platform"],
            )
            if batch_key in seen_in_batch:
                is_duplicate = True
            seen_in_batch.add(batch_key)
        results.append(
            {
                "row_number": idx,
                "raw": raw,
                "cleaned": cleaned,
                "errors": errors,
                "is_duplicate": is_duplicate,
                "valid": len(errors) == 0,
            }
        )
    return results


def import_rows(validated_rows: list[dict], skip_duplicates: bool = False) -> dict:
    imported = 0
    skipped_invalid = 0
    skipped_duplicate = 0
    for row in validated_rows:
        if not row["valid"]:
            skipped_invalid += 1
            continue
        if skip_duplicates and row["is_duplicate"]:
            skipped_duplicate += 1
            continue
        create_order(row["cleaned"])
        imported += 1
    return {"imported": imported, "skipped_invalid": skipped_invalid, "skipped_duplicate": skipped_duplicate}
